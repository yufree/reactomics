#!/usr/bin/env python3
"""Generate an editable PowerPoint for the reactomics paradigm-shift poster figures."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.oxml import parse_xml
from pptx.util import Inches, Pt, Emu

OUT = Path(__file__).resolve().parents[1] / "assets" / "poster"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(0x1A, 0x36, 0x5D)
TEAL = RGBColor(0x2B, 0x6C, 0xB0)
CORAL = RGBColor(0xC0, 0x56, 0x21)
GREEN = RGBColor(0x27, 0x67, 0x49)
PURPLE = RGBColor(0x6B, 0x46, 0xC1)
MUTED = RGBColor(0x4A, 0x55, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF0, 0xF7, 0xFC)
LIGHT = RGBColor(0xED, 0xF2, 0xF7)
EDGE = RGBColor(0x2C, 0x52, 0x82)


def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _set_line(shape, color, width_pt=1.5):
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def _no_line(shape):
    shape.line.fill.background()


def _textbox(slide, left, top, width, height, text, *, size=18, bold=False, color=NAVY,
             align=PP_ALIGN.CENTER, italic=False, font_name="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font_name
    return box


def _round_rect(slide, left, top, width, height, fill=WHITE, line=EDGE, line_pt=1.5):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _set_fill(shape, fill)
    if line is None:
        _no_line(shape)
    else:
        _set_line(shape, line, line_pt)
    # Soften corner a bit
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def _oval(slide, left, top, size, fill, line=NAVY, line_pt=2, label="", label_color=WHITE, label_size=16):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    _set_fill(shape, fill)
    _set_line(shape, line, line_pt)
    if label:
        tf = shape.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.size = Pt(label_size)
        run.font.bold = True
        run.font.color.rgb = label_color
        run.font.name = "Calibri"
        shape.text_frame.paragraphs[0].space_before = Pt(0)
    return shape


def _connector(slide, x1, y1, x2, y2, color=CORAL, width_pt=2.5):
    """Straight connector between two points (EMU)."""
    cxn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cxn.line.color.rgb = color
    cxn.line.width = Pt(width_pt)
    return cxn


def _set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def slide_paradigm_shift(prs):
    """Main editable paradigm-shift comparison slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _set_slide_bg(slide, BG)

    # Title
    _textbox(slide, Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.55),
             "Paradigm Shift: From Metabolites to Reactions",
             size=32, bold=True, color=NAVY)
    _textbox(slide, Inches(0.4), Inches(0.7), Inches(12.5), Inches(0.4),
             "Reactomics treats chemical reactions as first-class analytes in untargeted MS data",
             size=16, color=MUTED)

    # Left panel
    _round_rect(slide, Inches(0.35), Inches(1.25), Inches(5.5), Inches(5.6), fill=WHITE, line=EDGE, line_pt=2)
    _textbox(slide, Inches(0.5), Inches(1.4), Inches(5.2), Inches(0.45),
             "Traditional Metabolomics", size=22, bold=True, color=NAVY)

    # Left nodes (independent metabolites)
    left_nodes = {
        "A": (1.0, 2.2),
        "B": (3.2, 2.0),
        "C": (1.6, 3.3),
        "D": (3.8, 3.1),
        "E": (1.1, 4.4),
        "F": (3.5, 4.3),
    }
    node_size = Inches(0.55)
    for lab, (x, y) in left_nodes.items():
        _oval(slide, Inches(x), Inches(y), node_size, TEAL, NAVY, 2, lab, WHITE, 18)

    # Left captions
    _round_rect(slide, Inches(0.55), Inches(5.15), Inches(5.1), Inches(1.5), fill=LIGHT, line=None)
    _textbox(slide, Inches(0.6), Inches(5.25), Inches(5.0), Inches(0.35),
             "Metabolites as independent entities", size=16, italic=True, color=MUTED)
    _textbox(slide, Inches(0.6), Inches(5.6), Inches(5.0), Inches(0.4),
             'Question: "Which metabolites differ?"', size=18, bold=True, color=CORAL)
    _textbox(slide, Inches(0.6), Inches(6.05), Inches(5.0), Inches(0.35),
             "Requires compound annotation", size=15, color=MUTED)

    # SHIFT arrow
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(6.05), Inches(3.55), Inches(1.15), Inches(0.55)
    )
    _set_fill(arrow, CORAL)
    _no_line(arrow)
    _textbox(slide, Inches(6.05), Inches(3.15), Inches(1.15), Inches(0.35),
             "SHIFT", size=16, bold=True, color=CORAL)

    # Right panel
    _round_rect(slide, Inches(7.4), Inches(1.25), Inches(5.5), Inches(5.6), fill=WHITE, line=EDGE, line_pt=2)
    _textbox(slide, Inches(7.55), Inches(1.4), Inches(5.2), Inches(0.45),
             "Reactomics", size=22, bold=True, color=GREEN)

    right_nodes = {
        "A": (8.05, 2.2),
        "B": (10.25, 2.0),
        "C": (8.65, 3.3),
        "D": (10.85, 3.1),
        "E": (8.15, 4.4),
        "F": (10.55, 4.3),
    }
    # Draw edges first (behind nodes)
    edges = [
        ("A", "B", "CH₂  14.0157", CORAL),
        ("A", "C", "O  15.9949", TEAL),
        ("C", "E", "H₂  2.0157", GREEN),
        ("B", "D", "H₂O  18.0106", PURPLE),
        ("C", "D", "CH₂  14.0157", CORAL),
        ("E", "F", "O  15.9949", TEAL),
    ]
    centers = {}
    for lab, (x, y) in right_nodes.items():
        centers[lab] = (
            Inches(x) + node_size / 2,
            Inches(y) + node_size / 2,
        )

    for a, b, lab, col in edges:
        x1, y1 = centers[a]
        x2, y2 = centers[b]
        _connector(slide, x1, y1, x2, y2, col, 2.75)
        # Edge label near midpoint
        mx = (x1 + x2) / 2 - Inches(0.55)
        my = (y1 + y2) / 2 - Inches(0.18)
        _textbox(slide, mx, my, Inches(1.2), Inches(0.35), lab, size=12, bold=True, color=col)

    for lab, (x, y) in right_nodes.items():
        _oval(slide, Inches(x), Inches(y), node_size, WHITE, NAVY, 2, lab, NAVY, 18)

    # Right captions
    _round_rect(slide, Inches(7.6), Inches(5.15), Inches(5.1), Inches(1.5), fill=LIGHT, line=None)
    _textbox(slide, Inches(7.65), Inches(5.25), Inches(5.0), Inches(0.35),
             "Reactions (PMDs) as first-class objects", size=16, italic=True, color=MUTED)
    _textbox(slide, Inches(7.65), Inches(5.6), Inches(5.0), Inches(0.4),
             'Question: "Which reactions differ?"', size=18, bold=True, color=GREEN)
    _textbox(slide, Inches(7.65), Inches(6.05), Inches(5.0), Inches(0.35),
             "No compound ID required", size=15, color=MUTED)

    # Footer
    _textbox(slide, Inches(0.4), Inches(7.0), Inches(12.5), Inches(0.3),
             "PMD-based reactomics · Communications Chemistry 2020 · pmd R package",
             size=12, color=MUTED)


def slide_unit_of_analysis(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, BG)

    _textbox(slide, Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.5),
             "The Unit of Analysis Flips", size=30, bold=True, color=NAVY)

    # BEFORE
    _round_rect(slide, Inches(0.5), Inches(1.1), Inches(5.2), Inches(5.5), fill=WHITE, line=TEAL, line_pt=2.5)
    _textbox(slide, Inches(0.6), Inches(1.25), Inches(5.0), Inches(0.35),
             "BEFORE", size=14, bold=True, color=MUTED)
    _textbox(slide, Inches(0.6), Inches(1.6), Inches(5.0), Inches(0.45),
             "Unit = Molecule", size=26, bold=True, color=TEAL)

    for i, name in enumerate(["Glucose", "Alanine", "ATP", "…"]):
        y = 2.3 + i * 0.7
        card = _round_rect(slide, Inches(1.4), Inches(y), Inches(3.4), Inches(0.55), fill=RGBColor(0xEB, 0xF8, 0xFF), line=TEAL)
        _textbox(slide, Inches(1.4), Inches(y + 0.08), Inches(3.4), Inches(0.4),
                 name, size=18, bold=True, color=NAVY)

    _textbox(slide, Inches(0.7), Inches(5.6), Inches(4.8), Inches(0.7),
             "Reactions = secondary\n(edges between named nodes)", size=14, color=MUTED)

    # Arrow
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(6.0), Inches(3.5), Inches(1.2), Inches(0.55)
    )
    _set_fill(arrow, CORAL)
    _no_line(arrow)
    _textbox(slide, Inches(5.9), Inches(3.05), Inches(1.4), Inches(0.35),
             "REACTOMICS", size=12, bold=True, color=CORAL)

    # AFTER
    _round_rect(slide, Inches(7.5), Inches(1.1), Inches(5.2), Inches(5.5), fill=WHITE, line=GREEN, line_pt=2.5)
    _textbox(slide, Inches(7.6), Inches(1.25), Inches(5.0), Inches(0.35),
             "AFTER", size=14, bold=True, color=MUTED)
    _textbox(slide, Inches(7.6), Inches(1.6), Inches(5.0), Inches(0.45),
             "Unit = Reaction", size=26, bold=True, color=GREEN)

    reactions = [
        ("Oxidation  +O", TEAL),
        ("Methylation  +CH₂", CORAL),
        ("Phosphorylation  +HPO₃", PURPLE),
        ("…", MUTED),
    ]
    for i, (name, col) in enumerate(reactions):
        y = 2.3 + i * 0.7
        card = _round_rect(slide, Inches(8.4), Inches(y), Inches(3.4), Inches(0.55),
                           fill=RGBColor(0xF0, 0xFF, 0xF4), line=col, line_pt=2)
        _textbox(slide, Inches(8.4), Inches(y + 0.08), Inches(3.4), Inches(0.4),
                 name, size=16, bold=True, color=col)

    _textbox(slide, Inches(7.7), Inches(5.6), Inches(4.8), Inches(0.7),
             "Molecules = endpoints\n(may remain unnamed)", size=14, color=MUTED)

    _textbox(slide, Inches(0.5), Inches(6.85), Inches(12.3), Inches(0.4),
             "If endpoints must be named before a reaction is counted, the workflow has not left traditional metabolomics.",
             size=13, italic=True, color=MUTED)


def slide_pmd_detector(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, BG)

    _textbox(slide, Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.5),
             "Mass Spectrometry as a Chemical Reaction Detector", size=28, bold=True, color=NAVY)
    _textbox(slide, Inches(0.4), Inches(0.7), Inches(12.5), Inches(0.35),
             "A fixed paired mass distance (PMD) between two ions encodes a specific biotransformation",
             size=15, color=MUTED)

    # Spectrum panel
    _round_rect(slide, Inches(0.4), Inches(1.25), Inches(6.6), Inches(5.2), fill=WHITE, line=EDGE, line_pt=2)
    _textbox(slide, Inches(0.55), Inches(1.4), Inches(6.3), Inches(0.4),
             "High-resolution MS spectrum", size=18, bold=True, color=NAVY)

    # Baseline
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.9), Inches(5.5), Inches(6.4), Inches(5.5)
    )
    line.line.color.rgb = NAVY
    line.line.width = Pt(2)

    # Peaks as thin rectangles
    peaks = [
        (1.4, 1.2, "m/z₁"),
        (3.0, 2.6, "m/z₂  (substrate)"),
        (4.8, 2.0, "m/z₃  (product)"),
    ]
    for x, h, lab in peaks:
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(5.5 - h), Inches(0.12), Inches(h)
        )
        _set_fill(bar, TEAL)
        _no_line(bar)
        _textbox(slide, Inches(x - 0.55), Inches(5.5 - h - 0.45), Inches(1.3), Inches(0.4),
                 lab, size=12, bold=True, color=NAVY)

    # PMD double arrow (approximate with left+right arrows + label)
    _textbox(slide, Inches(2.9), Inches(5.65), Inches(2.2), Inches(0.35),
             "◀——  PMD = |m/z₃ − m/z₂|  ——▶", size=13, bold=True, color=CORAL)
    _textbox(slide, Inches(3.2), Inches(6.05), Inches(1.5), Inches(0.3),
             "m/z →", size=14, color=MUTED)

    # Reaction table
    _textbox(slide, Inches(7.4), Inches(1.35), Inches(5.3), Inches(0.4),
             "PMD  →  Reaction", size=20, bold=True, color=NAVY)

    rows = [
        ("2.0157 Da", "H₂", "Reduction", GREEN),
        ("14.0157 Da", "CH₂", "Methylation", CORAL),
        ("15.9949 Da", "O", "Oxidation", TEAL),
        ("18.0106 Da", "H₂O", "Hydration", PURPLE),
        ("79.9663 Da", "HPO₃", "Phosphorylation", NAVY),
        ("162.0528 Da", "C₆H₁₀O₅", "Glycosylation", RGBColor(0x9B, 0x2C, 0x2C)),
    ]
    for i, (pmd, formula, name, col) in enumerate(rows):
        y = 1.9 + i * 0.7
        card = _round_rect(slide, Inches(7.4), Inches(y), Inches(5.3), Inches(0.58), fill=WHITE, line=col, line_pt=1.75)
        _textbox(slide, Inches(7.55), Inches(y + 0.1), Inches(1.7), Inches(0.4),
                 pmd, size=14, bold=True, color=col, align=PP_ALIGN.LEFT)
        _textbox(slide, Inches(9.3), Inches(y + 0.1), Inches(1.4), Inches(0.4),
                 formula, size=14, bold=True, color=NAVY, align=PP_ALIGN.LEFT)
        _textbox(slide, Inches(10.6), Inches(y + 0.1), Inches(1.9), Inches(0.4),
                 name, size=14, color=MUTED, align=PP_ALIGN.RIGHT)

    _textbox(slide, Inches(0.4), Inches(6.7), Inches(12.5), Inches(0.35),
             "Core insight: fixed Δm between ions = chemical reaction linking them",
             size=15, bold=True, color=NAVY)


def slide_quantification(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, BG)

    _textbox(slide, Inches(0.3), Inches(0.15), Inches(12.7), Inches(0.45),
             "What Makes Reactomics an Omics: Reactions Are the Analytes", size=26, bold=True, color=NAVY)
    _textbox(slide, Inches(0.3), Inches(0.6), Inches(12.7), Inches(0.35),
             "Network construction alone is not enough — quantification at the reaction layer is the decisive shift",
             size=14, color=MUTED)

    steps = [
        (0.35, "1. Peak list", "Untargeted LC-MS\nm/z + intensity\n(no IDs needed)", TEAL),
        (4.55, "2. PMD pairs", "Match Δm to curated\nreactions (O, CH₂, …)\nacross all peak pairs", CORAL),
        (8.75, "3. Reaction matrix", "getreact(): reaction × sample\nactivity scores for stats\n& comparison", GREEN),
    ]
    for x, title, body, col in steps:
        _round_rect(slide, Inches(x), Inches(1.1), Inches(3.7), Inches(2.0), fill=WHITE, line=col, line_pt=2.5)
        _textbox(slide, Inches(x + 0.1), Inches(1.25), Inches(3.5), Inches(0.4),
                 title, size=18, bold=True, color=col)
        _textbox(slide, Inches(x + 0.15), Inches(1.75), Inches(3.4), Inches(1.2),
                 body, size=14, color=MUTED)

    # arrows between steps
    for x in (4.05, 8.25):
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(1.9), Inches(0.4), Inches(0.35))
        _set_fill(arr, NAVY)
        _no_line(arr)

    _textbox(slide, Inches(0.3), Inches(3.35), Inches(12.7), Inches(0.4),
             "Two regulation regimes → two quantitative readouts", size=18, bold=True, color=NAVY)

    # Static
    _round_rect(slide, Inches(0.35), Inches(3.85), Inches(6.0), Inches(2.9), fill=WHITE, line=TEAL, line_pt=2.5)
    _textbox(slide, Inches(0.5), Inches(4.0), Inches(5.7), Inches(0.4),
             "Static PMD", size=20, bold=True, color=TEAL)
    _textbox(slide, Inches(0.5), Inches(4.45), Inches(5.7), Inches(0.7),
             "Substrate & product co-vary\n(stable conversion efficiency)", size=14, color=MUTED)
    _textbox(slide, Inches(0.5), Inches(5.4), Inches(5.7), Inches(0.9),
             "Quantify: intensity sum (S + P)\n⇒ upstream / downstream control", size=16, bold=True, color=NAVY)

    # Dynamic
    _round_rect(slide, Inches(6.7), Inches(3.85), Inches(6.0), Inches(2.9), fill=WHITE, line=CORAL, line_pt=2.5)
    _textbox(slide, Inches(6.85), Inches(4.0), Inches(5.7), Inches(0.4),
             "Dynamic PMD", size=20, bold=True, color=CORAL)
    _textbox(slide, Inches(6.85), Inches(4.45), Inches(5.7), Inches(0.7),
             "Substrate & product diverge\n(enzyme activity changes)", size=14, color=MUTED)
    _textbox(slide, Inches(6.85), Inches(5.4), Inches(5.7), Inches(0.9),
             "Quantify: ratio (stable / variable)\n⇒ enzyme-level control", size=16, bold=True, color=NAVY)


def slide_consequences(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, BG)

    _textbox(slide, Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.5),
             "Consequences of Treating Reactions as First-Class Objects", size=26, bold=True, color=NAVY)

    cards = [
        (0.4, 1.1, "Robust to annotation gaps",
         "PMDs are computed for any peak pair — annotated or not. Missing IDs no longer block biological inference.",
         TEAL),
        (6.9, 1.1, "Chemically interpretable",
         "Δm = 2.0157 → reduction\nΔm = 14.0157 → methylation\nThe network is readable without a lookup table.",
         CORAL),
        (0.4, 4.1, "Comparable across systems",
         "Reaction networks can be compared across conditions, species, and sample types in ways metabolite lists cannot.",
         GREEN),
        (6.9, 4.1, "Links to pathway knowledge",
         "Connects to KEGG / HMDB / Reactome when IDs exist, yet remains usable when annotation is incomplete.",
         PURPLE),
    ]
    for x, y, title, body, col in cards:
        _round_rect(slide, Inches(x), Inches(y), Inches(5.9), Inches(2.6), fill=WHITE, line=col, line_pt=2.5)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(5.9), Inches(0.12))
        _set_fill(bar, col)
        _no_line(bar)
        _textbox(slide, Inches(x + 0.2), Inches(y + 0.35), Inches(5.5), Inches(0.5),
                 title, size=20, bold=True, color=col)
        _textbox(slide, Inches(x + 0.25), Inches(y + 1.0), Inches(5.4), Inches(1.3),
                 body, size=15, color=MUTED, align=PP_ALIGN.LEFT)


def slide_notes(prs):
    """Editing tips for the user."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_slide_bg(slide, WHITE)
    _textbox(slide, Inches(0.5), Inches(0.4), Inches(12), Inches(0.5),
             "How to edit this deck", size=28, bold=True, color=NAVY, align=PP_ALIGN.LEFT)
    tips = (
        "• Every element is a native PowerPoint shape or text box — select and restyle freely.\n\n"
        "• Slide 1 = main paradigm-shift figure (best starting point for the poster).\n\n"
        "• Slide 2 = unit of analysis flip · Slide 3 = PMD as reaction detector\n"
        "  Slide 4 = reaction-level quantification · Slide 5 = practical consequences.\n\n"
        "• Recommended poster workflow: set slide size to your poster panel, enlarge fonts,\n"
        "  then File → Export → PDF / PNG, or copy shapes into your poster master.\n\n"
        "• Current slide size: widescreen 13.333 × 7.5 in. Change via Design → Slide Size.\n\n"
        "• Pair with assets/pmd-orbitals.png for the quantized-PMD / orbital analogy panel."
    )
    _textbox(slide, Inches(0.5), Inches(1.2), Inches(12), Inches(5.5),
             tips, size=18, color=MUTED, align=PP_ALIGN.LEFT)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide_paradigm_shift(prs)
    slide_unit_of_analysis(prs)
    slide_pmd_detector(prs)
    slide_quantification(prs)
    slide_consequences(prs)
    slide_notes(prs)

    out = OUT / "reactomics_paradigm_shift.pptx"
    prs.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
